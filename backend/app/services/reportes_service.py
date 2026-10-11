"""FASE 12: reportes empresariales acotados por empresa, sin SQL dinámico."""
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from io import BytesIO, StringIO
import csv
import os
from html import escape
from uuid import uuid4
from sqlalchemy import select, func
from app.models.venta import Venta
from app.models.venta_detalle import VentaDetalle
from app.models.producto import Producto
from app.models.cliente import Cliente
from app.models.reporte_generado import ReporteGenerado
from app.services.dashboard_comercial_service import _filtro_ventas, _money
from app.services.venta_service import _validar_operador, VentaServiceError

TIPOS = {'VENTAS', 'ESTADISTICO', 'INVENTARIO', 'CLIENTES', 'RENDIMIENTO'}
FORMATOS = {'CSV', 'XLSX', 'PDF'}
ALMACEN = Path(os.getenv('SALESIA_REPORT_DIR', 'reports/generated')).resolve()

class ReporteError(ValueError):
    def __init__(self, mensaje, status_code=422):
        super().__init__(mensaje)
        self.status_code = status_code


def _tenant(db, user_id, permission):
    try:
        _, empresa = _validar_operador(db, usuario_id=user_id, permiso=permission)
    except VentaServiceError as exc:
        raise ReporteError('No autorizado para reportes.', 403) from exc
    return empresa.id


def _datos(db, empresa_id, tipo, inicio, fin):
    filtros = _filtro_ventas(empresa_id, inicio, fin)
    if tipo == 'VENTAS':
        consulta = (select(Venta.id, Venta.codigo, Venta.created_at, Venta.estado_pago, Venta.total)
            .where(*filtros).order_by(Venta.created_at, Venta.id).limit(50001))
        encabezados = ['id', 'codigo', 'fecha_utc', 'estado_pago', 'total_pen']
        filas = [[id_, codigo, fecha.isoformat(), estado_pago, _money(total)] for id_,codigo,fecha,estado_pago,total in db.execute(consulta)]
    elif tipo == 'INVENTARIO':  # Reporte de PRODUCTOS VENDIDOS (no stock físico)
        consulta = (select(Producto.id, Producto.nombre, func.sum(VentaDetalle.cantidad), func.sum(VentaDetalle.total))
            .join(VentaDetalle, (VentaDetalle.producto_id == Producto.id) & (VentaDetalle.empresa_id == Producto.empresa_id))
            .join(Venta, (Venta.id == VentaDetalle.venta_id) & (Venta.empresa_id == VentaDetalle.empresa_id))
            .where(*filtros).group_by(Producto.id, Producto.nombre).order_by(Producto.id).limit(50001))
        encabezados = ['producto_id', 'producto', 'unidades_vendidas', 'ventas_pen']
        filas = [[i, n, str(c), _money(t)] for i,n,c,t in db.execute(consulta)]
    elif tipo == 'CLIENTES':
        consulta = (select(Cliente.id, Cliente.nombre, func.count(Venta.id), func.sum(Venta.total))
            .join(Venta, (Venta.cliente_id == Cliente.id) & (Venta.empresa_id == Cliente.empresa_id))
            .where(*filtros).group_by(Cliente.id, Cliente.nombre).order_by(Cliente.id).limit(50001))
        encabezados = ['cliente_id', 'cliente', 'ventas_confirmadas', 'total_pen']
        filas = [[i, n, int(c), _money(t)] for i,n,c,t in db.execute(consulta)]
    elif tipo == 'RENDIMIENTO':
        consulta = (select(Venta.usuario_id, func.count(Venta.id), func.sum(Venta.total))
            .where(*filtros).group_by(Venta.usuario_id).order_by(Venta.usuario_id).limit(50001))
        encabezados = ['vendedor_usuario_id', 'ventas_confirmadas', 'total_pen']
        filas = [[i, int(c), _money(t)] for i,c,t in db.execute(consulta)]
    else:
        # Estadísticas descriptivas reproducibles de ventas confirmadas.
        cantidad, total, minimo, maximo = db.execute(select(func.count(Venta.id), func.coalesce(func.sum(Venta.total),0),
              func.min(Venta.total), func.max(Venta.total)).where(*filtros)).one()
        encabezados = ['indicador', 'valor']
        filas = [['ventas_confirmadas', int(cantidad)], ['total_pen', _money(total)],
                 ['ticket_promedio_pen', _money(Decimal(str(total))/cantidad) if cantidad else ''],
                 ['venta_minima_pen', _money(minimo) if minimo is not None else ''],
                 ['venta_maxima_pen', _money(maximo) if maximo is not None else '']]
    if len(filas) > 50000:
        raise ReporteError('Límite de 50000 filas excedido; reduzca las fechas.')
    return encabezados, filas


def _csv(encabezados, filas):
    out = StringIO(newline='')
    writer = csv.writer(out)
    # Prevención de inyección de fórmulas al abrir CSV en programas de hojas de cálculo.
    def seguro(x):
        s = str(x) if x is not None else ''
        return "'" + s if s.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')) else s
    writer.writerow(encabezados)
    writer.writerows([[seguro(v) for v in fila] for fila in filas])
    return ('\ufeff' + out.getvalue()).encode('utf-8')


def _xlsx(encabezados, filas):
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    sheet = wb.active
    sheet.title = 'Reporte'
    sheet.append([str(c) for c in encabezados])
    for fila in filas:
        valores = []
        for valor in fila:
            texto = str(valor) if valor is not None else ''
            # Almacenar datos como texto para evitar ejecución de fórmulas.
            if texto.lstrip().startswith(('=', '+', '-', '@', '\t', '\r', '\n')):
                texto = "'" + texto
            valores.append(texto)
        sheet.append(valores)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='253B59')
        cell.alignment = Alignment(wrap_text=True)
    sheet.freeze_panes = 'A2'
    if encabezados:
        sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(encabezados, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(len(str(header)) + 3, 15), 40)
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()

def _pdf(encabezados, filas, titulo, inicio, fin):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import A4, landscape
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=landscape(A4), leftMargin=30, rightMargin=30)
    styles = getSampleStyleSheet()
    elementos = [Paragraph(escape(titulo), styles['Title']), Paragraph(f'{inicio} a {fin} | UTC | PEN', styles['Normal']), Spacer(1,12)]
    # Acotar la cantidad de filas en PDF por memoria y longitud del documento.
    if len(filas) > 2000:
        raise ReporteError("PDF admite hasta 2000 filas; exporte CSV o XLSX para períodos mayores.")
    # Truncar solo la presentación de texto en celdas, no los datos de las exportaciones.
    ancho = (landscape(A4)[0]-60)/len(encabezados)
    tabla = [[Paragraph(escape(str(c)), styles['BodyText']) for c in encabezados]]
    for fila in filas:
        tabla.append([Paragraph(escape(str(v)[:140] if v is not None else ''), styles['BodyText']) for v in fila])
    if len(tabla) == 1: tabla.append(['Sin resultados'] + ['']*(len(encabezados)-1))
    t = Table(tabla, colWidths=[ancho]*len(encabezados), repeatRows=1)
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e3eaf4')),
         ('VALIGN',(0,0),(-1,-1),'TOP'), ('BOTTOMPADDING',(0,0),(-1,-1),7),
         ('LINEBELOW',(0,0),(-1,0),.5,colors.grey)]))
    elementos.append(t)
    doc.build(elementos)
    return stream.getvalue()


def _html(encabezados, filas, titulo, inicio, fin):
    cab = ''.join(f'<th>{escape(str(c))}</th>' for c in encabezados)
    body = ''.join('<tr>'+''.join(f'<td>{escape(str(v))}</td>' for v in fila)+'</tr>' for fila in filas)
    return ('<!doctype html><html lang="es"><head><meta charset="utf-8"><title>Reporte</title>'
            '<style>body{font-family:Arial;margin:32px}table{border-collapse:collapse;width:100%}'
            'td,th{border:1px solid #bbb;padding:7px;text-align:left}th{background:#e3eaf4}'
            '@media print{button{display:none}}</style></head><body>'
            f'<button onclick="window.print()">Imprimir</button><h1>{escape(titulo)}</h1>'
            f'<p>{inicio} a {fin} | UTC | PEN</p><table><thead><tr>{cab}</tr></thead>'
            f'<tbody>{body}</tbody></table></body></html>').encode('utf-8')


def previsualizar(db, usuario_id, tipo, inicio, fin, limite=100):
    eid = _tenant(db, usuario_id, 'reportes.ver')
    validar(tipo, inicio, fin)
    h, rows = _datos(db, eid, tipo, inicio, fin)
    return {'tipo': tipo, 'fecha_inicio': inicio, 'fecha_fin': fin, 'columnas': h,
            'total_filas': len(rows), 'filas': rows[:limite], 'moneda': 'PEN', 'zona_horaria': 'UTC'}


def validar(tipo, inicio, fin):
    if tipo not in TIPOS: raise ReporteError('Tipo de reporte no admitido.')
    if fin < inicio or (fin-inicio).days > 366: raise ReporteError('Intervalo inválido o mayor a 367 días.')


def generar(db, usuario_id, tipo, formato, inicio, fin):
    eid = _tenant(db, usuario_id, 'reportes.generar')
    validar(tipo, inicio, fin)
    if formato not in FORMATOS: raise ReporteError('Formato no admitido.')
    headers, rows = _datos(db, eid, tipo, inicio, fin)
    titulo = f'SalesIA - {tipo}'
    try:
        contenido = {'CSV': lambda: _csv(headers, rows), 'XLSX': lambda: _xlsx(headers, rows),
                     'PDF': lambda: _pdf(headers, rows, titulo, inicio, fin)}[formato]()
    except ImportError as exc:
        raise ReporteError(f'Dependencia de exportación no instalada: {exc}', 503) from exc
    if not contenido: raise ReporteError('Archivo de reporte vacío.', 500)
    # Ruta privada generada internamente. Nunca aceptar rutas enviadas por clientes.
    relativo = f'{eid}/{uuid4().hex}.{formato.lower()}'
    destino = (ALMACEN / relativo).resolve()
    if not destino.is_relative_to(ALMACEN): raise ReporteError('Ruta no autorizada.', 500)
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_bytes(contenido)
    registro = ReporteGenerado(empresa_id=eid, solicitado_por_id=usuario_id,
        nombre=f'{tipo} {inicio} a {fin}', tipo_reporte=tipo, formato=formato,
        estado='COMPLETADO', parametros={'fecha_inicio':inicio.isoformat(),'fecha_fin':fin.isoformat()},
        ruta_archivo=relativo, tamano_bytes=len(contenido), fecha_generacion=datetime.now(timezone.utc))
    db.add(registro)
    try:
        db.commit()
        db.refresh(registro)
    except Exception:
        db.rollback()
        destino.unlink(missing_ok=True)
        raise
    return registro


def listar(db, usuario_id, limite=50):
    eid = _tenant(db, usuario_id, 'reportes.ver')
    return db.execute(select(ReporteGenerado).where(ReporteGenerado.empresa_id == eid)
          .order_by(ReporteGenerado.created_at.desc(), ReporteGenerado.id.desc()).limit(limite)).scalars().all()


def consultar(db, usuario_id, reporte_id, permiso):
    eid = _tenant(db, usuario_id, permiso)
    registro = db.execute(select(ReporteGenerado).where(ReporteGenerado.id == reporte_id,
        ReporteGenerado.empresa_id == eid)).scalar_one_or_none()
    if registro is None: raise ReporteError('Reporte no encontrado.', 404)
    return registro


def descargar(db, usuario_id, reporte_id):
    r = consultar(db, usuario_id, reporte_id, 'reportes.descargar')
    if r.estado != 'COMPLETADO' or not r.ruta_archivo: raise ReporteError('Reporte no disponible.', 409)
    archivo = (ALMACEN / r.ruta_archivo).resolve()
    if not archivo.is_relative_to(ALMACEN / str(r.empresa_id)):
        raise ReporteError('Ruta de reporte inválida.', 500)
    if not archivo.is_file(): raise ReporteError('Archivo no disponible en este servidor.', 404)
    return r, archivo


def imprimible(db, usuario_id, tipo, inicio, fin):
    eid = _tenant(db, usuario_id, 'reportes.ver')
    validar(tipo, inicio, fin)
    h, rows = _datos(db, eid, tipo, inicio, fin)
    return _html(h, rows, f'SalesIA - {tipo}', inicio, fin)
